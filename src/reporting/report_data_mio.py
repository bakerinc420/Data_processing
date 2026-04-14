import os
import io
import numpy as np
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from datetime import datetime

# Импорты ваших модулей
from src.database.provider.reader import EMGReader
from src.core.algorithms.time_domain import TimeDomainProcessor
from src.core.algorithms.frequency_domain import FrequencyDomainProcessor
from src.utils.time_visualizer import EMGVisualizer
from config import REPORTS_DIR, FS

class MioReportGenerator:
    def __init__(self):
        self.reader = EMGReader()
        self.time_proc = TimeDomainProcessor()
        self.freq_proc = FrequencyDomainProcessor()
        self.viz_t = EMGVisualizer()

    def _get_plot_stream(self, fig):
        """Превращает график в поток байтов для Word."""
        stream = io.BytesIO()
        fig.savefig(stream, format='png', bbox_inches='tight')
        plt.close(fig)
        stream.seek(0)
        return stream

    def create_combined_mio_report(self, p_id):
        """Создает полный отчет по базе data_mio для одного пациента."""
        doc = Document()
        doc.add_heading(f'ОТЧЕТ ЭМГ: ПАЦИЕНТ №{p_id}', 0)
        
        table_name = "data_mio"
        full_signal = self.reader.get_signal(p_id, table_name) #
        
        if full_signal is None or len(full_signal) == 0:
            print(f"❌ Данные для пациента {p_id} не найдены.")
            return

        for trial in range(1, 9):
            trial_idx = trial - 1
            doc.add_heading(f'АНАЛИЗ ПРОБЫ №{trial}', level=1)
            
            # Нарезка исходного сигнала (24000 точек)
            sample = full_signal[trial_idx * 24000 : trial * 24000]
            if len(sample) == 0 or np.all(sample == 0):
                continue

            # 1. Поиск границ активности (синхронизация с main_time.py)
            start_act, end_act = self.time_proc.find_activity_borders(sample)
            
            # 2. Выделение активной зоны для синхронных расчетов
            if start_act is not None and end_act is not None:
                active_zone = sample[start_act:end_act]
            else:
                active_zone = sample
                start_act, end_act = 0, len(sample) - 1

            # 3. Подготовка данных для алгоритмов
            active_zone_centered = active_zone - np.mean(active_zone)
            stats = self.time_proc.calculate_stats(active_zone_centered, segment_size=100) #

            # --- РАЗДЕЛ 1: ВРЕМЕННОЙ АНАЛИЗ ---
            doc.add_heading('1. Временной анализ', level=2)
            
            p_time = doc.add_paragraph()
            p_time.add_run(f"• Средняя амплитуда (M+STD): {stats['custom_avg']:.2f} мкВ\n")
            p_time.add_run(f"• Медианная амплитуда: {stats['median_amp']:.2f} мкВ\n")
            p_time.add_run(f"• RMS (мощность сигнала): {stats['rms']:.2f} мкВ\n")
            p_time.add_run(f"• iEMG (интегральная активность): {stats['iemg']:.2f} мкВ*с\n")
            p_time.add_run(f"• ZCR (частота пересеч. нуля): {stats['zcr']:.2f} Гц\n")
            p_time.add_run(f"• Амплитуда от пика до пика (P-P): {stats['peak_to_peak']:.2f} мкВ\n")
            p_time.add_run(f"• Максимальная амплитуда: {stats['max_amp']:.2f} мкВ\n")
            p_time.add_run(f"• Минимальная амплитуда: {stats['min_amp']:.2f} мкВ")

            # ГРАФИК 1: Окно 1 (Добавлены точки Max/Min)
            fig1 = plt.figure(figsize=(10, 4))
            plt.plot(sample, color='gray', alpha=0.5, label='ЭМГ сигнал')
            plt.axvline(x=start_act, color='green', linestyle='--', label=f'Start: {start_act}')
            plt.axvline(x=end_act, color='red', linestyle='--', label=f'End: {end_act}')
            
            # Добавление точек максимума и минимума
            plt.scatter(np.argmax(sample), np.max(sample), color='darkred', zorder=5, 
                        label=f'Max: {np.max(sample):.2f}')
            plt.scatter(np.argmin(sample), np.min(sample), color='blue', zorder=5, 
                        label=f'Min: {np.min(sample):.2f}')
            
            plt.title(f"Окно 1: Сигнал и Активность (Проба {trial})")
            plt.legend(loc='upper right')
            plt.grid(True, alpha=0.3)
            doc.add_picture(self._get_plot_stream(fig1), width=Inches(6))

            # ГРАФИК 2: Окно 2 (Огибающая и Линейная аппроксимация)
            fig2, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 8))
            envelope = self.viz_t._create_peak_envelope(sample)
            
            k2_db, k3_db = self.reader.get_specific_borders(p_id, trial_idx) #
            safe_k3 = k2_db if (k2_db and k2_db != 0) else start_act
            safe_k4 = k3_db if (k3_db and k3_db != 0) else end_act

            y3_env, y4_env = envelope[safe_k3], envelope[safe_k4]
            x_range = np.arange(len(sample))
            if safe_k4 != safe_k3:
                slope_t = (y4_env - y3_env) / (safe_k4 - safe_k3)
                lin_approx = slope_t * (x_range - safe_k3) + y3_env
            else:
                lin_approx = np.full_like(x_range, y3_env)

            ax1.plot(sample, color='gray', alpha=0.15)
            ax1.plot(envelope, color='red', linewidth=2, label='Огибающая')
            ax1.set_title("Огибающая (Peak Envelope)")
            
            ax2.plot(sample, color='gray', alpha=0.2)
            ax2.plot(envelope, color='red', alpha=0.3) 
            ax2.plot(lin_approx, color='blue', linewidth=3, linestyle='--', label='Линия тренда (k3-k4)')
            ax2.set_title("Линейная апроксимация (k3-k4)")
            ax2.grid(True, alpha=0.3)
            ax2.legend()
            
            plt.tight_layout()
            doc.add_picture(self._get_plot_stream(fig2), width=Inches(6))

            # --- РАЗДЕЛ 2: ЧАСТОТНЫЙ АНАЛИЗ ---
            doc.add_heading('2. Частотный анализ', level=2)
            f_res = self.freq_proc.get_frequency_report(active_zone_centered) #
            
            p_freq = doc.add_paragraph()
            p_freq.add_run(f"• Медианная частота (MDF): {f_res['median_frequency']:.2f} Гц\n")
            p_freq.add_run(f"• Средняя частота (MNF): {f_res['mean_frequency']:.2f} Гц\n")
            p_freq.add_run(f"• Пиковая частота: {f_res['peak_frequency']:.2f} Гц\n")
            p_freq.add_run(f"• Коэффициент H/L (баланс): {f_res['hl_ratio']:.3f}\n")
            p_freq.add_run(f"• Спектральная энтропия: {f_res['entropy']:.3f}\n")
            p_freq.add_run(f"• Утомление (Slope): {f_res['fatigue_slope']:.2f} Гц/сек\n")
            p_freq.add_run(f"• Падение MDF за пробу: {f_res['fatigue_drop_pct']:.1f} %")

            # График 3: PSD (с линиями MDF и MNF)
            fig3 = plt.figure(figsize=(10, 5))
            plt.plot(f_res['freqs'], f_res['psd'], color='black', linewidth=1.2)
            plt.axvline(x=f_res['median_frequency'], color='green', linestyle='--', linewidth=2, label='MDF')
            plt.axvline(x=f_res['mean_frequency'], color='blue', linestyle=':', linewidth=2, label='MNF')
            plt.xlim(0, 500)
            plt.title("Спектральная плотность мощности (PSD)")
            plt.legend()
            plt.grid(True, linestyle='--', alpha=0.4)
            doc.add_picture(self._get_plot_stream(fig3), width=Inches(6))

            # График 4: Спектрограмма (синхронная шкала дБ)
            fig4 = plt.figure(figsize=(10, 5))
            from scipy.signal import spectrogram
            f, t, Sxx = spectrogram(sample, fs=FS, nperseg=1024, noverlap=900)
            pc = plt.pcolormesh(t, f, 10 * np.log10(Sxx + 1e-10), 
                                shading='gouraud', cmap='magma', 
                                vmin=-25, vmax=15)
            plt.ylim(20, 500)
            plt.title(f"Спектрограмма: data_mio | Пациент {p_id}, Проба {trial}")
            plt.colorbar(pc, label="Мощность (дБ)")
            doc.add_picture(self._get_plot_stream(fig4), width=Inches(6))
            
            doc.add_page_break()

        if not os.path.exists(REPORTS_DIR):
            os.makedirs(REPORTS_DIR)
        doc.save(os.path.join(REPORTS_DIR, f"Patient_{p_id}_Full_Analysis.docx"))
        print(f"✅ Отчет для пациента {p_id} успешно сформирован.")

if __name__ == "__main__":
    MioReportGenerator().create_combined_mio_report(p_id=1)